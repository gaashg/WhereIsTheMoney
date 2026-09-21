# WhereIsTheMoney — Copyright (C) 2026 GaashG
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version. See <https://www.gnu.org/licenses/>.
#
# Additional terms under Section 7(b): see the NOTICE file.
import logging
from dataclasses import fields
from typing import Any, cast

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter

from utils import config
from models.expense import Expense

logger = logging.getLogger(__name__)

# The dataclass decides which columns are taken, and in which order.
FIELDS = [field.name for field in fields(Expense)]
# A row saying nothing in any of these describes no purchase at all: it is
# either empty, or a totals row, whose sum sits in the billing amount alone.
PURCHASE_COLUMNS = ["shop", "purchase_amount", "purchase_date"]


def read_file(file_path: str) -> list[Expense] | None:
    wanted = set(config.get_value("columns_names", "categories_table_structure"))
    columns_renaming = config.get_value("column_names_renaming", "categories_table_structure")

    workbook = load_workbook(file_path, data_only=True)
    failures = {}
    # search for the appropriate table in the workbook.
    for worksheet in workbook.worksheets:
        for table in worksheet.tables.values():
            # if wanted <= {c.name for c in table.tableColumns}:
            if set(columns_renaming.keys()) <= {c.name for c in table.tableColumns}:
                # A table ref always carries both columns and rows, unlike the
                # partial refs ("B:F") that range_boundaries also accepts.
                min_col, min_row, max_col, max_row = cast(
                    tuple[int, int, int, int], range_boundaries(table.ref))
                if None in (min_col, min_row, max_col, max_row):
                    raise ValueError(f"Unusable table range: {table.ref}")
                df = pd.read_excel(
                    file_path,
                    sheet_name=worksheet.title,
                    header=0,
                    # first row of the table is its header
                    skiprows=min_row - 1,
                    # rows above the table
                    nrows=max_row - min_row,
                    # data rows, header excluded
                    usecols=f"{get_column_letter(min_col)}:{get_column_letter(max_col)}",
                )
                df.rename(columns=columns_renaming, inplace=True)
                # An empty row or a totals row, not a purchase.
                empty_rows = df[PURCHASE_COLUMNS].isna().all(axis=1)
                for index in df.index[empty_rows]:
                    logger.warning("Row %s holds no purchase, dropping it: %s",
                                   index, df.loc[index].to_dict())
                df = df[~empty_rows]
                df["purchase_date"] = _dated(df)
                # Empty cells arrive as NaN floats, which say nothing here.
                wanted_columns = df[FIELDS].astype(object)
                wanted_columns = wanted_columns.where(pd.notna(wanted_columns), None)
                # The keys are the column names, so they are strings, as ** needs.
                records = cast(list[dict[str, Any]], wanted_columns.to_dict("records"))
                return [Expense(**row) for row in records]
            else:
                failures[table.ref] = wanted - {c.name for c in table.tableColumns}

    logger.error(f"No suitable table data was found for file {file_path}. "
                 f"Inconsistencies: {failures}")
    return None


def _dated(df: pd.DataFrame) -> pd.Series:
    """The purchase dates, as real dates, with every gap filled.

    Excel hands the same column over in three shapes: real dates, serial
    numbers counted from 1899-12-30, and dd/MM/yyyy text. Whatever is left
    unreadable, and whatever was empty, is dated to the first day of the month
    most of the purchases belong to.
    """
    column = df["purchase_date"]
    # A plain number is a serial: read as text it would become 1970-01-01.
    # As objects, real dates are not numbers; as datetime64 they would be.
    numbers = pd.to_numeric(column.astype(object), errors="coerce")
    serials = pd.to_datetime(numbers, unit="D", origin="1899-12-30", errors="coerce")
    dates = pd.to_datetime(column.where(numbers.isna()), dayfirst=True,
                           errors="coerce").fillna(serials)

    months = dates.dt.to_period("M").mode()
    if months.empty:
        raise ValueError("No purchase date in the table could be read")

    first_of_month = months[0].to_timestamp()
    for index in df.index[dates.isna()]:
        logger.warning("Entry %s (%s) has no purchase date, its date was set to %s",
                       index, df.at[index, "shop"],
                       first_of_month.strftime("%d/%m/%Y"))

    return dates.fillna(first_of_month)
