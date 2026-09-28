# WhereIsTheMoney — Copyright (C) 2026 GaashG
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version. See <https://www.gnu.org/licenses/>.
#
# Additional terms under Section 7(b): see the NOTICE file.

"""Tests for writers.expenses_analyzer_file_writer.

Every test points the config at its own folder under tmp_path, so the real
output folder is never written to. The template itself is the real one: it is
what the tests are about.
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import cast

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from exceptions import ConfigError
from models.expense import Expense
from utils import config
from writers import expenses_analyzer_file_writer as writer

CONFIG = """
[categorized_expenses]
categorized_expenses_file_location = '{folder}'
categorized_expenses_file_prefix = "expensesFor%d-%d.xlsx"
"""

HEADER_ROW = 15


def expense(shop="SUPER", date=datetime(2024, 11, 9), purchase_amount=121.71,
            billing_amount=121.71, detail=None, category1="מזון",
            category2="מכולת", category3=None) -> Expense:
    return Expense(purchase_date=date, shop=shop, category1=category1,
                   category2=category2, category3=category3,
                   purchase_amount=purchase_amount, billing_amount=billing_amount,
                   extra_detail=detail)


@pytest.fixture
def output_folder(tmp_path, monkeypatch):
    """Point the config at a folder of this test's own, and return it."""

    def configure(folder: Path, text: str = CONFIG) -> Path:
        config_file = tmp_path / "config.toml"
        config_file.write_text(text.format(folder=folder), encoding="utf-8")
        monkeypatch.setenv(config.CONFIG_PATH_ENV_VAR, str(config_file))
        config._load_toml.cache_clear()
        return folder

    folder = configure(tmp_path / "categorized")
    folder.mkdir()
    yield folder
    config._load_toml.cache_clear()


@pytest.fixture
def written(output_folder):
    """Return a function that writes expenses and reads the result back."""

    def write(expenses, month=11, year=2024):
        writer.write_categorized_expenses(month, year, expenses)
        path = output_folder / f"expensesFor{year}-{month:02d}.xlsx"
        workbook = load_workbook(path)
        worksheet = next(sheet for sheet in workbook.worksheets
                         if writer.EXPENSES_TABLE in sheet.tables)
        return workbook, worksheet

    return write


def rows_of(worksheet) -> list[list]:
    """The data rows of the expenses table, the totals row left out."""
    first_column, header, last_column, totals = cast(
        tuple[int, int, int, int],
        range_boundaries(worksheet.tables[writer.EXPENSES_TABLE].ref))
    return [[cell.value for cell in row] for row in
            worksheet.iter_rows(min_row=header + 1, max_row=totals - 1,
                                min_col=first_column, max_col=last_column)]


# --- positive cases: the file itself ----------------------------------------


@pytest.mark.parametrize(
    "month, year, name",
    [
        (11, 2024, "expensesFor2024-11.xlsx"),
        (3, 2024, "expensesFor2024-03.xlsx"),
        (1, 2026, "expensesFor2026-01.xlsx"),
        (12, 2025, "expensesFor2025-12.xlsx"),
    ],
    ids=["two digit month", "padded month", "january", "december"],
)
def test_the_year_and_the_month_make_the_file_name(output_folder, month, year, name):
    writer.write_categorized_expenses(month, year, [expense()])

    assert [path.name for path in output_folder.iterdir()] == [name]


def test_a_folder_that_does_not_exist_yet_is_created(tmp_path, monkeypatch):
    folder = tmp_path / "budget" / "categorized"
    config_file = tmp_path / "config.toml"
    config_file.write_text(CONFIG.format(folder=folder), encoding="utf-8")
    monkeypatch.setenv(config.CONFIG_PATH_ENV_VAR, str(config_file))
    config._load_toml.cache_clear()

    writer.write_categorized_expenses(11, 2024, [expense()])

    assert (folder / "expensesFor2024-11.xlsx").exists()
    config._load_toml.cache_clear()


def test_an_earlier_file_of_the_same_month_is_replaced(written):
    written([expense("FIRST"), expense("SECOND")])
    _, worksheet = written([expense("ONLY")])

    assert [row[1] for row in rows_of(worksheet)] == ["ONLY"]


def test_the_path_is_logged(output_folder, caplog):
    with caplog.at_level(logging.INFO):
        writer.write_categorized_expenses(11, 2024, [expense()])

    assert "expensesFor2024-11.xlsx" in caplog.text


# --- positive cases: the expenses table -------------------------------------


def test_an_expense_fills_a_row_of_the_table(written):
    _, worksheet = written([expense("חצי חינם", datetime(2024, 11, 3), 121.71, 40.57,
                                    "תשלום 1 מתוך 3", "אוכל", "מכולת", "שבועי")])

    assert rows_of(worksheet) == [[datetime(2024, 11, 3), "חצי חינם", 121.71, 40.57,
                                   "תשלום 1 מתוך 3", "אוכל", "מכולת", "שבועי"]]


def test_every_expense_gets_a_row_of_its_own(written):
    _, worksheet = written([expense("A"), expense("B"), expense("C")])

    assert [row[1] for row in rows_of(worksheet)] == ["A", "B", "C"]


def test_the_first_expense_goes_right_under_the_header(written):
    _, worksheet = written([expense("FIRST")])

    assert worksheet.cell(row=HEADER_ROW + 1, column=2).value == "FIRST"


def test_empty_fields_are_left_empty(written):
    _, worksheet = written([expense(detail=None, category3=None)])

    row = rows_of(worksheet)[0]
    assert (row[4], row[7]) == (None, None)


def test_the_table_covers_every_expense(written):
    _, worksheet = written([expense() for _ in range(5)])

    # The header row, five expenses and the totals row.
    assert worksheet.tables[writer.EXPENSES_TABLE].ref == "A15:H21"


def test_the_totals_row_sits_below_the_last_expense(written):
    _, worksheet = written([expense(), expense()])

    assert "SUBTOTAL" in worksheet.cell(row=HEADER_ROW + 3, column=4).value


def test_a_long_list_of_expenses_stretches_the_table(written):
    _, worksheet = written([expense() for _ in range(120)])

    assert len(rows_of(worksheet)) == 120
    assert worksheet.tables[writer.EXPENSES_TABLE].ref == "A15:H136"


def test_the_dates_and_the_amounts_keep_their_formats(written):
    _, worksheet = written([expense() for _ in range(3)])

    last = HEADER_ROW + 3
    assert "d/m/yy" in worksheet.cell(row=last, column=1).number_format
    assert worksheet.cell(row=last, column=3).number_format == "#,##0.00"
    assert worksheet.cell(row=last, column=4).number_format == "#,##0.00"


def test_dates_are_written_as_dates_and_amounts_as_numbers(written):
    _, worksheet = written([expense()])

    row = rows_of(worksheet)[0]
    assert isinstance(row[0], datetime)
    assert isinstance(row[2], float)


# --- positive cases: what the template brings along --------------------------


def test_the_other_two_tables_come_with_the_template(written):
    _, worksheet = written([expense()])

    assert set(worksheet.tables) == {"טבלה1", "טבלה2", "טבלה3"}


def test_the_income_table_is_left_empty_for_its_owner(written):
    _, worksheet = written([expense()])

    assert worksheet["A5"].value is not None  # the row's label
    assert worksheet["B5"].value is None      # the amount, filled in by hand


def test_the_formulas_of_the_template_are_kept(written):
    _, worksheet = written([expense()])

    assert worksheet["B8"].value == "=SUM(B5:B7)"
    assert "SUBTOTAL" in worksheet["L5"].value


def test_the_pivot_table_survives_and_is_refreshed_when_opened(written):
    workbook, _ = written([expense()])

    pivots = [pivot for sheet in workbook.worksheets for pivot in sheet._pivots]
    assert len(pivots) == 1
    assert pivots[0].cache.refreshOnLoad is True
    assert pivots[0].cache.cacheSource.worksheetSource.name == writer.EXPENSES_TABLE


def test_the_sheet_is_still_right_to_left(written):
    _, worksheet = written([expense()])

    assert worksheet.sheet_view.rightToLeft is True


# --- negative cases ---------------------------------------------------------


def test_no_expenses_still_give_a_file_with_the_table(written):
    _, worksheet = written([])

    assert rows_of(worksheet) == [[None] * 8]
    assert worksheet.tables[writer.EXPENSES_TABLE].ref == "A15:H17"


@pytest.mark.parametrize(
    "key",
    ["categorized_expenses_file_location", "categorized_expenses_file_prefix"],
)
def test_a_missing_config_key_is_reported(tmp_path, monkeypatch, key):
    kept = "\n".join(line for line in CONFIG.splitlines() if not line.startswith(key))
    config_file = tmp_path / "config.toml"
    config_file.write_text(kept.format(folder=tmp_path), encoding="utf-8")
    monkeypatch.setenv(config.CONFIG_PATH_ENV_VAR, str(config_file))
    config._load_toml.cache_clear()

    with pytest.raises(ConfigError, match=key):
        writer.write_categorized_expenses(11, 2024, [expense()])

    config._load_toml.cache_clear()


def test_a_template_that_is_not_there_is_reported(output_folder, monkeypatch, tmp_path):
    monkeypatch.setattr(writer, "TEMPLATE_PATH", tmp_path / "no_such_template.xlsx")

    with pytest.raises(FileNotFoundError):
        writer.write_categorized_expenses(11, 2024, [expense()])


def test_a_template_without_the_expenses_table_is_reported(output_folder, monkeypatch,
                                                           tmp_path):
    monkeypatch.setattr(writer, "EXPENSES_TABLE", "no such table")

    with pytest.raises(ValueError, match="no such table"):
        writer.write_categorized_expenses(11, 2024, [expense()])


def test_a_value_excel_cannot_hold_is_reported(output_folder):
    with pytest.raises(ValueError):
        writer.write_categorized_expenses(11, 2024, [expense(detail=["a", "list"])])


def test_a_folder_in_the_place_of_the_output_file_is_reported(output_folder):
    (output_folder / "expensesFor2024-11.xlsx").mkdir()

    with pytest.raises(OSError):
        writer.write_categorized_expenses(11, 2024, [expense()])
