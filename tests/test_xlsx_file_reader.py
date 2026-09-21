# WhereIsTheMoney — Copyright (C) 2026 GaashG
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version. See <https://www.gnu.org/licenses/>.
#
# Additional terms under Section 7(b): see the NOTICE file.

"""Tests for file_readers.xlsx_file_reader.

Every test builds its own workbook, with real Excel tables, under tmp_path, and
points the config at a copy of the table structure, so neither the real data
nor the real config.toml are involved.
"""

import logging
import zipfile
from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from file_readers import xlsx_file_reader
from models.expense import Expense
from utils import config

CONFIG = """
[categories_table_structure]
columns_names = ["תאריך רכישה", "שם בית עסק", "סכום עסקה"]
column_names_renaming = {"תאריך רכישה" = "purchase_date", "שם בית עסק" = "shop", "סכום עסקה" = "purchase_amount", "סכום חיוב" = "billing_amount", "תחום 1" = "category1", "תחום 2" = "category2", "תחום 3" = "category3"}
"""

HEADERS = ["תאריך רכישה", "שם בית עסק", "סכום עסקה", "סכום חיוב", "פירוט נוסף",
           "תחום 1", "תחום 2", "תחום 3"]


def purchase(date, shop="NETFLIX", amount=10.0, billing=10.0, note=None,
             category1="תקשורת", category2=None, category3=None) -> list:
    """One table row, in the order of HEADERS."""
    return [date, shop, amount, billing, note, category1, category2, category3]


def add_table(worksheet, rows, *, row=1, column=1, headers=HEADERS, name="Table1"):
    """Write the headers and rows at (row, column) and make them an Excel table."""
    for offset, values in enumerate([headers, *rows]):
        for shift, value in enumerate(values):
            worksheet.cell(row=row + offset, column=column + shift, value=value)

    first = f"{get_column_letter(column)}{row}"
    last = f"{get_column_letter(column + len(headers) - 1)}{row + len(rows)}"
    worksheet.add_table(Table(displayName=name, ref=f"{first}:{last}"))


@pytest.fixture(autouse=True)
def table_structure(tmp_path, monkeypatch):
    """Make the reader use the test's copy of the table structure."""
    config_file = tmp_path / "config.toml"
    config_file.write_text(CONFIG, encoding="utf-8")
    monkeypatch.setenv(config.CONFIG_PATH_ENV_VAR, str(config_file))
    config._load_toml.cache_clear()
    yield
    config._load_toml.cache_clear()


@pytest.fixture
def save(tmp_path):
    """Return a function that saves a workbook and gives back its path."""

    def write(workbook: Workbook, name: str = "expenses.xlsx") -> str:
        path = tmp_path / name
        workbook.save(path)
        return str(path)

    return write


@pytest.fixture
def one_table(save):
    """Return a function that saves a workbook holding a single table of rows."""

    def write(*rows) -> str:
        workbook = Workbook()
        add_table(workbook.active, list(rows))
        return save(workbook)

    return write


# --- positive cases: reading ------------------------------------------------


def test_rows_become_expenses(one_table):
    path = one_table(purchase(datetime(2024, 10, 9), "NETFLIX", category1="תקשורת",
                              category2="טלוויזיה", category3="נטפליקס"))

    assert xlsx_file_reader.read_file(path) == [
        Expense(purchase_date=datetime(2024, 10, 9), shop="NETFLIX",
                category1="תקשורת", category2="טלוויזיה", category3="נטפליקס")]


def test_empty_categories_become_none(one_table):
    path = one_table(purchase(datetime(2024, 10, 9), category2=None, category3=None))

    expense = xlsx_file_reader.read_file(path)[0]
    assert expense.category2 is None
    assert expense.category3 is None


def test_rows_keep_their_order(one_table):
    path = one_table(purchase(datetime(2024, 10, 3), "A"),
                     purchase(datetime(2024, 10, 1), "B"),
                     purchase(datetime(2024, 10, 2), "C"))

    assert [e.shop for e in xlsx_file_reader.read_file(path)] == ["A", "B", "C"]


def test_identical_rows_are_all_kept(one_table):
    """Two identical purchases are two purchases, not one."""
    same = purchase(datetime(2024, 10, 9), "SUPER")
    path = one_table(same, same, same)

    assert len(xlsx_file_reader.read_file(path)) == 3


def test_table_away_from_the_corner(save):
    """Rows and columns above and left of the table are not part of it."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "a title the bank put here"
    worksheet["A2"] = 12345
    add_table(worksheet, [purchase(datetime(2024, 10, 9), "SHOP")], row=6, column=3)

    assert [e.shop for e in xlsx_file_reader.read_file(save(workbook))] == ["SHOP"]


def test_table_is_found_on_any_sheet_among_other_tables(save):
    """The sheet name, the sheet order and other tables do not matter."""
    workbook = Workbook()
    workbook.active.title = "summary"
    add_table(workbook.active, [["salary", 100]], headers=["הכנסות", "סכום"],
              name="Incomes")
    worksheet = workbook.create_sheet("גיליון1")
    add_table(worksheet, [["balance", 5]], headers=["יתרה", "סכום"], name="Balance")
    add_table(worksheet, [purchase(datetime(2024, 10, 9), "SHOP")], row=5,
              name="Purchases")

    assert [e.shop for e in xlsx_file_reader.read_file(save(workbook))] == ["SHOP"]


def test_only_the_first_matching_table_is_read(save):
    """A file holds one relevant table, so reading stops once it is processed."""
    workbook = Workbook()
    add_table(workbook.active, [purchase(datetime(2024, 10, 9), "FIRST")], name="T1")
    add_table(workbook.create_sheet("second"),
              [purchase(datetime(2024, 10, 9), "SECOND")], name="T2")

    assert [e.shop for e in xlsx_file_reader.read_file(save(workbook))] == ["FIRST"]


# --- positive cases: dates --------------------------------------------------


def test_text_dates_are_read_day_first(one_table):
    """09/10/2024 is the 9th of October, not the 10th of September."""
    path = one_table(purchase("09/10/2024"))

    assert xlsx_file_reader.read_file(path)[0].purchase_date == datetime(2024, 10, 9)


def test_serial_numbers_are_read_as_excel_dates(one_table):
    """Excel counts days from 1899-12-30: 45606 is 2024-11-10."""
    path = one_table(purchase(45606))

    assert xlsx_file_reader.read_file(path)[0].purchase_date == datetime(2024, 11, 10)


def test_mixed_date_shapes_in_one_column(one_table):
    path = one_table(purchase(datetime(2024, 10, 2)),
                     purchase("18/10/2024"),
                     purchase(45587))

    dates = [e.purchase_date for e in xlsx_file_reader.read_file(path)]
    assert dates == [datetime(2024, 10, 2), datetime(2024, 10, 18), datetime(2024, 10, 22)]


def test_missing_date_gets_the_first_of_the_most_common_month(one_table):
    """Older purchases (instalments, late billing) do not decide the month."""
    path = one_table(purchase(datetime(2024, 7, 12)),
                     purchase(datetime(2024, 10, 5)),
                     purchase(datetime(2024, 10, 20)),
                     purchase(None, "CASH"))

    assert xlsx_file_reader.read_file(path)[-1].purchase_date == datetime(2024, 10, 1)


def test_unreadable_date_is_treated_like_a_missing_one(one_table):
    path = one_table(purchase(datetime(2024, 10, 5)), purchase("not a date", "ODD"))

    assert xlsx_file_reader.read_file(path)[-1].purchase_date == datetime(2024, 10, 1)


def test_a_defaulted_date_is_logged(one_table, caplog):
    path = one_table(purchase(datetime(2024, 10, 5)), purchase(None, "CASH"))

    with caplog.at_level(logging.WARNING):
        xlsx_file_reader.read_file(path)

    assert "Entry 1 (CASH) has no purchase date, its date was set to 01/10/2024" \
        in caplog.text


def test_no_date_is_ever_left_empty(one_table):
    path = one_table(purchase(datetime(2024, 10, 5)), purchase(None),
                     purchase("garbage"), purchase(45606))

    assert all(e.purchase_date is not None for e in xlsx_file_reader.read_file(path))


# --- positive cases: rows that are not purchases ----------------------------


def test_totals_row_is_dropped(one_table):
    """Only the billing amount is filled: that is the table's total."""
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"),
                     purchase(None, shop=None, amount=None, billing=6023.15,
                              category1=None))

    assert [e.shop for e in xlsx_file_reader.read_file(path)] == ["SHOP"]


def test_empty_row_is_dropped(one_table):
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"),
                     purchase(None, shop=None, amount=None, billing=None,
                              category1=None))

    assert [e.shop for e in xlsx_file_reader.read_file(path)] == ["SHOP"]


def test_a_dropped_row_is_logged(one_table, caplog):
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"),
                     purchase(None, shop=None, amount=None, billing=6023.15,
                              category1=None))

    with caplog.at_level(logging.WARNING):
        xlsx_file_reader.read_file(path)

    assert "Row 1 holds no purchase, dropping it" in caplog.text
    assert "6023.15" in caplog.text


def test_a_dropped_row_does_not_get_a_default_date(one_table, caplog):
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"),
                     purchase(None, shop=None, amount=None, billing=6023.15,
                              category1=None))

    with caplog.at_level(logging.WARNING):
        xlsx_file_reader.read_file(path)

    assert "has no purchase date" not in caplog.text


@pytest.mark.parametrize(
    "row",
    [
        purchase(None, shop="CASH", amount=None, billing=None),
        purchase(None, shop=None, amount=400.0, billing=None),
        purchase(datetime(2024, 10, 5), shop=None, amount=None, billing=None),
    ],
    ids=["shop only", "purchase amount only", "date only"],
)
def test_a_row_with_any_purchase_detail_is_kept(one_table, row):
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"), row)

    assert len(xlsx_file_reader.read_file(path)) == 2


def test_a_row_without_a_shop_is_kept_with_shop_none(one_table):
    """Dropping shopless expenses is the categories parser's job, not the reader's."""
    path = one_table(purchase(datetime(2024, 10, 5), "SHOP"),
                     purchase(datetime(2024, 10, 6), shop=None, amount=400.0))

    assert xlsx_file_reader.read_file(path)[1].shop is None


# --- negative cases ---------------------------------------------------------


def test_workbook_without_tables_returns_none(save, caplog):
    workbook = Workbook()
    workbook.active["A1"] = "no tables here"

    with caplog.at_level(logging.ERROR):
        assert xlsx_file_reader.read_file(save(workbook)) is None

    assert "No suitable table data was found" in caplog.text


def test_table_missing_a_needed_column_returns_none(save, caplog):
    """Without "תחום 3" the table cannot become Expense objects."""
    workbook = Workbook()
    add_table(workbook.active, [purchase(datetime(2024, 10, 9))[:-1]],
              headers=HEADERS[:-1])

    with caplog.at_level(logging.ERROR):
        assert xlsx_file_reader.read_file(save(workbook)) is None

    assert "No suitable table data was found" in caplog.text


def test_the_error_names_the_missing_columns(save, caplog):
    workbook = Workbook()
    add_table(workbook.active, [["salary", 100]], headers=["הכנסות", "סכום"])

    with caplog.at_level(logging.ERROR):
        xlsx_file_reader.read_file(save(workbook))

    assert "שם בית עסק" in caplog.text


def test_table_without_any_readable_date_raises(one_table):
    path = one_table(purchase(None, "A"), purchase("garbage", "B"))

    with pytest.raises(ValueError, match="No purchase date in the table could be read"):
        xlsx_file_reader.read_file(path)


def test_table_with_only_a_header_raises(save):
    """No rows means no month to date anything by."""
    workbook = Workbook()
    worksheet = workbook.active
    for column, header in enumerate(HEADERS, start=1):
        worksheet.cell(row=1, column=column, value=header)
    worksheet.add_table(Table(displayName="Table1", ref="A1:H2"))

    with pytest.raises(ValueError):
        xlsx_file_reader.read_file(save(workbook))


def test_missing_file_raises_file_not_found(tmp_path):
    with pytest.raises(FileNotFoundError):
        xlsx_file_reader.read_file(str(tmp_path / "absent.xlsx"))


def test_file_that_is_not_really_xlsx_raises(tmp_path):
    fake = tmp_path / "fake.xlsx"
    fake.write_text("just text wearing an .xlsx suffix", encoding="utf-8")

    with pytest.raises(zipfile.BadZipFile):
        xlsx_file_reader.read_file(str(fake))
