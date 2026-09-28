# WhereIsTheMoney — Copyright (C) 2026 GaashG
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU Affero General Public License as published
# by the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version. See <https://www.gnu.org/licenses/>.
#
# Additional terms under Section 7(b): see the NOTICE file.

"""Writes the categorized expenses into a monthly budget workbook.

The workbook is not built from nothing: openpyxl cannot create a pivot table,
only keep one it read. So every file starts as a copy of the template, which
holds the three tables, their formulas and the pivot table, and only the rows
of the expenses table are filled in.
"""

import logging
from pathlib import Path
from typing import cast

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from models.expense import Expense
from utils import config

logger = logging.getLogger(__name__)

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "resources" / "expenses_template.xlsx"
# The expenses table. The pivot table on the second sheet reads from it by name.
EXPENSES_TABLE = "טבלה2"
# The expense fields, in the order of the table's columns.
COLUMNS = ["purchase_date", "shop", "purchase_amount", "billing_amount",
           "extra_detail", "category1", "category2", "category3"]


def write_categorized_expenses(month: int, year: int,
                               categorized_expenses: list[Expense]):
    path = _output_path(month, year)
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet, table = _expenses_table(workbook)
    _fill(worksheet, table, categorized_expenses)
    path.parent.mkdir(parents=True, exist_ok=True)
    logger.info("Writing %s expenses into %s", len(categorized_expenses), path)
    # An earlier file for the same month is replaced.
    workbook.save(path)


def _output_path(month: int, year: int) -> Path:
    folder = config.get_value("categorized_expenses_file_location", "categorized_expenses")
    name_format = config.get_value("categorized_expenses_file_prefix", "categorized_expenses")
    # The first %d takes the year, the second the month, padded to two digits.
    file_name = name_format.replace("%d", str(year), 1).replace("%d", f"{month:02d}", 1)
    return Path(folder) / file_name


def _expenses_table(workbook: Workbook) -> tuple[Worksheet, Table]:
    """The template's expenses table, and the sheet it sits in."""
    for worksheet in workbook.worksheets:
        if EXPENSES_TABLE in worksheet.tables:
            return worksheet, worksheet.tables[EXPENSES_TABLE]

    raise ValueError(f"Template {TEMPLATE_PATH} holds no table named {EXPENSES_TABLE}")


def _fill(worksheet: Worksheet, table: Table, expenses: list[Expense]):
    """Put every expense in a row of its own, and move the totals row below them."""
    # A table ref always carries both columns and rows, unlike the partial refs
    # ("B:F") that range_boundaries also accepts.
    first_column, header_row, last_column, totals_row = cast(
        tuple[int, int, int, int], range_boundaries(table.ref))
    first_row = header_row + 1
    # The template's own row holds the formats the written values need.
    formats = [worksheet.cell(row=first_row, column=column).number_format
               for column in range(first_column, last_column + 1)]
    totals = {cell.column: cell.value for cell in worksheet[totals_row] if cell.value}
    for cell in worksheet[totals_row]:
        cell.value = None

    for offset, expense in enumerate(expenses):
        for index, field in enumerate(COLUMNS):
            cell = worksheet.cell(row=first_row + offset, column=first_column + index,
                                  value=getattr(expense, field))
            cell.number_format = formats[index]

    # The table keeps a row even with nothing to write, so the pivot has a source.
    new_totals_row = max(first_row + len(expenses), first_row + 1)
    for column, formula in totals.items():
        cell = worksheet.cell(row=new_totals_row, column=column, value=formula)
        cell.number_format = formats[column - first_column]

    table.ref = (f"{get_column_letter(first_column)}{header_row}:"
                 f"{get_column_letter(last_column)}{new_totals_row}")
